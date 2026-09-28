import openpyxl

wb = openpyxl.load_workbook(r"c:\Pratik_Bhuwad\VANA\GROUP2_THANE_CREEK\01_SOURCE_ARTIFACTS\SAKSHI\Scientific_Source_Registry_V0.1_FINAL_VERIFIED.xlsx")

with open(r"c:\Pratik_Bhuwad\VANA\GROUP2_THANE_CREEK\08_ANTIGRAVITY_ANALYSIS\sakshi_verified_dump.txt", "w", encoding="utf-8") as f:
    for name in wb.sheetnames:
        sheet = wb[name]
        f.write(f"\n=================== Sheet: {name} ({sheet.max_row} rows) ===================\n")
        for r in range(1, sheet.max_row + 1):
            row_vals = [str(cell.value) if cell.value is not None else "" for cell in sheet[r]]
            if any(row_vals):
                f.write(f"Row {r:02d}: {row_vals}\n")
print("Done writing sheets to sakshi_verified_dump.txt")
